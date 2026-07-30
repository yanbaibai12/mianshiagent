from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import quote


INTERNAL_KEYS = {
    "rag_references",
    "ats_report",
    "resume_evidence",
    "resume_index",
    "jd_alignment",
    "change_details",
}

SCORE_LABELS = {
    "technical_accuracy": "技术准确性",
    "project_understanding": "项目理解",
    "structure_clarity": "表达结构",
    "troubleshooting": "问题定位",
    "engineering_delivery": "工程落地",
    "reflection": "复盘能力",
    "completeness": "完整性",
    "logic": "逻辑清晰度",
    "consistency": "简历一致性",
    "conciseness": "表达精炼度",
    "depth": "技术/业务深度",
}


def _score_label(key: str) -> str:
    return SCORE_LABELS.get(key, key)

FIELD_LABELS = {
    "name": "姓名",
    "phone": "电话",
    "email": "邮箱",
    "job_intent": "求职意向",
    "school": "学校",
    "major": "专业",
    "degree": "学历",
    "time": "时间",
    "company": "公司",
    "role": "角色",
    "position": "岗位",
    "title": "标题",
    "description": "描述",
    "summary": "概述",
    "responsibilities": "职责",
    "highlights": "亮点",
    "achievements": "成果",
    "tech_stack": "技术栈",
    "interview_points": "可面试要点",
    "projects": "项目经历",
    "experience": "实习/工作经历",
    "education": "教育经历",
    "skills": "技能",
}


def safe_filename(value: str, suffix: str) -> str:
    cleaned = "".join(char if char.isalnum() or char in "-_（）() " else "_" for char in value).strip()
    cleaned = "_".join(cleaned.split())
    return f"{cleaned[:80] or 'export'}{suffix}"


def content_disposition(filename: str) -> str:
    ascii_fallback = "".join(char if ord(char) < 128 and char not in '";\\' else "_" for char in filename)
    ascii_fallback = ascii_fallback or "export"
    return f"attachment; filename=\"{ascii_fallback}\"; filename*=UTF-8''{quote(filename)}"


def _as_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, (int, float, bool)):
        return str(value)
    if isinstance(value, list):
        return "\n".join(_as_text(item) for item in value if _as_text(item))
    if isinstance(value, dict):
        lines = []
        for key, item in value.items():
            if key in INTERNAL_KEYS:
                continue
            text = _as_text(item)
            if text:
                lines.append(f"{FIELD_LABELS.get(str(key), str(key))}：{text}")
        return "\n".join(lines)
    return str(value)


def _items(value: Any) -> list[dict[str, Any]]:
    return [item for item in value if isinstance(item, dict)] if isinstance(value, list) else []


def _list_text(items: Any, limit: int | None = None) -> str:
    values = [str(item) for item in (items or []) if str(item).strip()] if isinstance(items, list) else []
    if limit:
        values = values[:limit]
    return "、".join(values)


def resume_to_docx_bytes(
    data: dict[str, Any],
    *,
    title: str,
    jd_text: str | None = None,
    ats_report: dict[str, Any] | None = None,
    change_details: list[Any] | None = None,
    job_context: dict[str, Any] | None = None,
) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading(title, level=0)
    personal = data.get("personal") if isinstance(data.get("personal"), dict) else {}
    name = personal.get("name") or "候选人"
    doc.add_heading(str(name), level=1)
    contact = " | ".join(str(personal.get(key) or "") for key in ["phone", "email", "job_intent"] if personal.get(key))
    if contact:
        doc.add_paragraph(contact)
    if job_context:
        doc.add_heading("岗位信息", level=1)
        company = job_context.get("company") or ""
        job_title = job_context.get("title") or ""
        if company or job_title:
            doc.add_paragraph(f"{company} {job_title}".strip())
    elif jd_text:
        doc.add_heading("岗位信息", level=1)
    if jd_text:
        doc.add_paragraph(jd_text[:1200])

    sections = [
        ("技能", data.get("skills")),
        ("教育经历", data.get("education")),
        ("实习/工作经历", data.get("experience")),
        ("项目经历", data.get("projects")),
        ("自我评价", data.get("summary")),
    ]
    for heading, value in sections:
        text = _as_text(value)
        if not text:
            continue
        doc.add_heading(heading, level=1)
        if isinstance(value, list) and all(isinstance(item, dict) for item in value):
            for item in _items(value):
                item_title = item.get("name") or item.get("company") or item.get("school") or item.get("title") or heading
                doc.add_heading(str(item_title), level=2)
                for line in _as_text(item).splitlines():
                    if line.strip():
                        doc.add_paragraph(line.strip(), style="List Bullet")
        elif isinstance(value, list):
            for item in value:
                doc.add_paragraph(str(item), style="List Bullet")
        else:
            doc.add_paragraph(text)

    ats_report = ats_report if isinstance(ats_report, dict) else data.get("ats_report") if isinstance(data.get("ats_report"), dict) else {}
    if ats_report:
        doc.add_heading("JD 匹配评分", level=1)
        doc.add_paragraph(f"总分：{ats_report.get('total_score', '-')}/100")
        for dimension in ats_report.get("dimensions") or []:
            if not isinstance(dimension, dict):
                continue
            doc.add_heading(str(dimension.get("name") or dimension.get("key") or "评分维度"), level=2)
            doc.add_paragraph(f"得分：{dimension.get('score', '-')}")
            if dimension.get("covered"):
                doc.add_paragraph("已覆盖：" + "、".join(map(str, dimension.get("covered") or [])))
            if dimension.get("missing"):
                doc.add_paragraph("缺口：" + "、".join(map(str, dimension.get("missing") or [])))
            if dimension.get("suggestion"):
                doc.add_paragraph("建议：" + str(dimension.get("suggestion")))
            for evidence in (dimension.get("evidence") or [])[:2]:
                if isinstance(evidence, dict) and evidence.get("excerpt"):
                    doc.add_paragraph(f"证据：{evidence.get('section_label') or evidence.get('section', '')} · {evidence.get('excerpt')}")
        advice = ats_report.get("delivery_advice") or []
        if advice:
            doc.add_heading("可投递建议", level=2)
            for item in advice:
                doc.add_paragraph(str(item), style="List Bullet")

    changes = change_details if isinstance(change_details, list) else data.get("change_details") if isinstance(data.get("change_details"), list) else []
    if changes:
        doc.add_heading("修改前后对比", level=1)
        table = doc.add_table(rows=1, cols=4)
        header_cells = table.rows[0].cells
        header_cells[0].text = "模块"
        header_cells[1].text = "修改前"
        header_cells[2].text = "修改后"
        header_cells[3].text = "改动原因"
        for item in changes[:30]:
            if not isinstance(item, dict):
                continue
            cells = table.add_row().cells
            cells[0].text = str(item.get("section") or item.get("module") or "")
            cells[1].text = str(item.get("before") or "")[:600]
            cells[2].text = str(item.get("after") or "")[:600]
            cells[3].text = str(item.get("reason") or item.get("evidence") or "")[:600]

    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def interview_report_to_markdown(report: Any) -> str:
    lines = [
        "# 面试总结报告\n",
        f"**总体评分**: {report.total_score}\n",
        "\n## 分项得分\n",
    ]
    for key, value in (report.dimension_scores or {}).items():
        lines.append(f"- {_score_label(str(key))}: {value}\n")
    lines.extend([f"\n## 表现总结\n\n{report.summary}\n", "\n## 薄弱点\n"])
    for i, point in enumerate(report.weak_points or [], 1):
        lines.append(f"{i}. {point}\n")
    lines.append("\n## 改进建议\n")
    for i, suggestion in enumerate(report.suggestions or [], 1):
        lines.append(f"{i}. {suggestion}\n")
    lines.append("\n## 答题详情\n")
    for q in report.questions:
        module = getattr(q, "module", "") or "resume"
        source = getattr(q, "source_section", "") or getattr(q, "point_title", "") or ""
        lines.append(f"\n### {q.question}\n")
        if module or source:
            lines.append(f"**模块**: {module}  \n**来源**: {source}\n")
        lines.append(f"**你的回答**: {q.user_answer or '未作答'}\n")
        if q.total_score is not None:
            lines.append(f"**评分**: {q.total_score}\n")
        if q.feedback:
            lines.append(f"**评语**: {q.feedback}\n")
        if q.refined_answer:
            lines.append(f"**精简答案**: {q.refined_answer}\n")
    return "".join(lines)


def interview_report_to_docx_bytes(report: Any) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("面试总结报告", level=0)
    doc.add_paragraph(f"总体评分：{report.total_score}")
    doc.add_heading("表现总结", level=1)
    doc.add_paragraph(report.summary or "")
    doc.add_heading("分项得分", level=1)
    for key, value in (report.dimension_scores or {}).items():
        doc.add_paragraph(f"{_score_label(str(key))}: {value}", style="List Bullet")
    doc.add_heading("薄弱点", level=1)
    for point in report.weak_points or []:
        doc.add_paragraph(str(point), style="List Bullet")
    doc.add_heading("改进建议", level=1)
    for suggestion in report.suggestions or []:
        doc.add_paragraph(str(suggestion), style="List Bullet")
    doc.add_heading("答题详情", level=1)
    for question in report.questions:
        doc.add_heading(question.question, level=2)
        module = getattr(question, "module", "") or "resume"
        source = getattr(question, "source_section", "") or getattr(question, "point_title", "") or ""
        doc.add_paragraph(f"模块：{module}；来源：{source}")
        doc.add_paragraph(f"你的回答：{question.user_answer or '未作答'}")
        if question.total_score is not None:
            doc.add_paragraph(f"评分：{question.total_score}")
        if question.feedback:
            doc.add_paragraph(f"评语：{question.feedback}")
        if question.refined_answer:
            doc.add_paragraph(f"精简答案：{question.refined_answer}")
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def _wrap_line(line: str, width: int = 46) -> list[str]:
    if len(line) <= width:
        return [line]
    return [line[index : index + width] for index in range(0, len(line), width)]


def text_to_pdf_bytes(text: str, *, title: str = "报告") -> bytes:
    import fitz

    doc = fitz.open()
    fontfile = None
    for candidate in [
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simsun.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
    ]:
        if candidate.exists():
            fontfile = str(candidate)
            break

    page = doc.new_page(width=595, height=842)
    y = 48
    lines: list[str] = []
    for raw_line in [title, "", *text.splitlines()]:
        lines.extend(_wrap_line(raw_line))
    for line in lines:
        if y > 790:
            page = doc.new_page(width=595, height=842)
            y = 48
        page.insert_text((48, y), line[:90], fontsize=11, fontfile=fontfile)
        y += 18
    return doc.tobytes()
